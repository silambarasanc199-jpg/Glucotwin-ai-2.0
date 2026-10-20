"""Build the editable presentation and print-ready PDF assets from local reports."""

import json
from pathlib import Path
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
REPORT = ROOT / "models" / "evaluation.json"
sys.path.insert(0, str(ROOT))


def slides_data() -> list[tuple[str, str, list[str]]]:
    report = json.loads(REPORT.read_text(encoding="utf-8")) if REPORT.exists() else None
    example_bullets = [
        "Show the actual stored profile and timestamps from the local API.",
        "Run the prediction in the dashboard and read the returned score without treating it as care advice.",
        "No numeric example is inserted if the generated dataset or trained model is unavailable.",
    ]
    data_file = ROOT / "data" / "synthetic_glucose.csv"
    if data_file.exists() and report:
        try:
            import pandas as pd
            from app.ml.features import FEATURE_COLUMNS, build_training_frame
            from app.services.prediction_service import load_bundle, make_prediction

            generated = pd.read_csv(data_file)
            bundle = load_bundle()
            training = build_training_frame(generated)
            candidates = training[(training.future_spike == 1) & (training.glucose_now < 180)].copy()
            if candidates.empty:
                candidates = training.copy()
            candidates["_score"] = bundle["pipeline"].predict_proba(candidates[FEATURE_COLUMNS])[:, 1]
            example_row = candidates.sort_values("_score").iloc[-1]
            source_id = str(example_row.patient_id)
            records = generated.loc[generated.patient_id == source_id].reset_index(drop=True)
            record_times = pd.to_datetime(records.timestamp, utc=True)
            peak_index = int(records.index[record_times == example_row.timestamp][0])
            profile = records.iloc[0]
            score = make_prediction({"patient_id": source_id, "bmi": float(profile.bmi),
                                     "hba1c": float(profile.hba1c), "is_synthetic": True},
                                    records.loc[max(0, peak_index - 31):peak_index].to_dict("records"),
                                    bundle)
            example_bullets = [
                f"Generated synthetic profile {source_id}; measured input at {score['latest_input_timestamp']} was {float(records.loc[peak_index, 'glucose_mg_dl']):.1f} mg/dL.",
                f"Actual saved model returned raw score {score['raw_model_score']:.4f}; uncalibrated, not an individual clinical probability.",
                "The simulated next-eight-row target is positive at this point. This historical synthetic example is stale and is not a current reading or care recommendation.",
            ]
        except (ImportError, OSError, KeyError, ValueError, FileNotFoundError):
            pass
    if report:
        test = report["test_metrics_at_threshold_0_5"]
        validation = report["validation_metrics_at_threshold_0_5"]
        results = [
            f"Validation PR-AUC: prior {validation['majority_baseline']['pr_auc']:.4f}; logistic {validation['logistic_regression']['pr_auc']:.4f}; Random Forest {validation['random_forest']['pr_auc']:.4f}.",
            f"Selected by validation PR-AUC: {report['selected_model']}.",
            f"Held-out test at fixed 0.5: precision {test['precision']:.4f}; recall {test['recall']:.4f}; F1 {test['f1']:.4f}.",
            f"PR-AUC {test['pr_auc']:.4f}; ROC-AUC {test['roc_auc']:.4f}; n={test['sample_count']:,} rows from {report['split']['patient_counts']['test']} held-out synthetic profiles.",
            f"Confusion matrix [[TN, FP], [FN, TP]]: {test['confusion_matrix_labels_0_1']}.",
            "These are synthetic-only results, not clinical validation; the test was not used for model selection.",
        ]
    else:
        results = ["Run `python -m app.ml.generate_synthetic_data` and `python -m app.ml.train` to produce actual results.",
                   "No performance numbers are fabricated or inserted when a report is missing."]
    return [
        ("GlucoTwin AI 2.0", "Personalized Diabetes Digital Twin · Happiest Health Digital Twin Challenge 2026", [
            "Project owner / team leader: SILAMBARASAN C", "Muthayammal Engineering College · Tamil Nadu, India",
            "Research prototype · all included demo data is synthetic", "Not a medical device; not for diagnosis or care decisions"]),
        ("The healthcare problem", "Glucose varies over time; retrospective review may not express a near-term individualized pattern.", [
            "Historical health context and current readings are often reviewed separately.", "A useful prototype should make time, data freshness, and missingness visible.",
            "This project demonstrates a technical workflow, not proven patient benefit."]),
        ("Proposed solution", "Link an anonymous profile to historical context and a dynamic observation timeline.", [
            "Store age, BMI, diabetes duration, and HbA1c only when available.", "Append timestamped glucose, heart rate, sleep, and step counts.",
            "Estimate a bounded future event and return input timestamp, warnings, and a safety disclaimer."]),
        ("Why a digital twin?", "An inspectable, updateable representation of a person's stored state.", [
            "Historical attributes and changing observations remain linked by anonymous patient ID.", "Clinicians/evaluators can inspect the records behind a model output.",
            "This is a data twin proof of concept, not a full physiological replica."]),
        ("Two data streams", "Static context + dynamic time-series observations.", [
            "Static: age, BMI, optional HbA1c and diabetes duration.", "Dynamic: timezone-aware glucose and optional heart rate, sleep, and activity.",
            "Synthetic trajectories are generated reproducibly and explicitly marked; no public or patient data is bundled."]),
        ("System architecture", "Validate → persist → engineer past-only features → infer → inspect.", [
            "FastAPI request schemas validate units, ranges, timestamps, and duplicate source records.", "SQLite stores anonymous profiles, readings, and predictions.",
            "A separate offline pipeline saves a trusted joblib artifact and held-out evaluation report.", "See docs/architecture_diagram.pdf and editable docs/architecture_diagram.svg."]),
        ("Dataset and ethics", "48,384 synthetic readings · 36 generated profiles · 14 days · 15-minute intervals.", [
            "Fixed seed 20260626; simplified meal pulses, circadian variation, noise, activity, sleep, and static attributes.",
            "Assumptions are software-test fixtures, not a clinical simulator.", "No external/public dataset was supplied or used; all demo records are labelled synthetic."]),
        ("Past-only feature engineering", "Use current and earlier observations as inputs; reserve future glucose for labels only.", [
            "Current glucose, 15- and 60-minute changes, 60-minute mean, and 120-minute variability.",
            "Heart rate, sleep, trailing steps, optional BMI/HbA1c, and time-of-day encoding.", "Training-fitted median imputation; irregular 15-minute training cadence is rejected."]),
        ("Target and training design", "Prototype label: any of the next eight 15-minute glucose readings reaches ≥180 mg/dL.", [
            "This is an explicit project target, not a universal medical standard.", "Patient-group holdout: 21 train, 7 validation, 8 test profiles.",
            "Prior baseline, Logistic Regression, and Random Forest compared; selected on validation PR-AUC.", "Training-only class weighting; no test-set tuning; default threshold 0.5 is reported."]),
        ("Model selection and results", "Actual generated results are reported only when the evaluation artifact exists.", results),
        ("Prediction workflow", "Latest readings → past-only feature row → model score → persisted response.", [
            "API returns patient ID, prediction time, 120-minute horizon, event definition, and model version.",
            "Response includes uncalibrated raw score, input timestamp, observed signals, global model ranking, quality warnings, and disclaimer.",
            "If training was not completed, inference returns a clear service error and training command."]),
        ("Dashboard walkthrough", "The browser dashboard reads patient, observations, score history, and evaluation from FastAPI.", [
            "Select backend-returned profiles; inspect measured-only glucose chart and source timeline.", "Add a manual observation without a fake sensor connection.",
            "Review stale/missing context, synthetic labels, score limitations, and backend evaluation."]),
        ("Synthetic patient example", "A generated synthetic profile; the prediction is an experimental model output.", [
            *example_bullets]),
        ("Limitations and next steps", "The current score does not establish safety, usefulness, or generalization.", [
            "Only synthetic data; simplified generator; one project-defined threshold; small number of held-out synthetic profiles.",
            "No external validation, calibration, confidence interval, subgroup analysis, clinical review, or prospective evaluation.",
            "No authentication or production security; SQLite is local demonstration storage.",
            "Next: licensed representative data, ethics/clinical review, external validation, calibration, subgroup evaluation, and security controls."]),
        ("Conclusion and references", "A reproducible full-stack technical proof of concept, not a clinical product.", [
            "Project documentation: README, model card, data dictionary, limitations, and 21-minute demo script.",
            "References: FastAPI documentation; scikit-learn documentation; SQLite documentation; Chart.js documentation.",
            "Team leader must verify submission rules, add only confirmed team details, and insert actual repository/video links.",
            "Thank you · Happiest Health Digital Twin Challenge 2026"]),
    ]


def add_powerpoint(slides: list[tuple[str, str, list[str]]]) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]
    screenshot = DOCS / "dashboard_synthetic.png"
    for index, (title, subtitle, bullets) in enumerate(slides):
        slide = presentation.slides.add_slide(blank)
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = RGBColor(20, 49, 72) if index == 0 else RGBColor(244, 248, 251)
        if index == 0:
            shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10.5), Inches(-1.25), Inches(4.3), Inches(4.3))
            shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor(30, 144, 157); shape.line.fill.background()
            title_box = slide.shapes.add_textbox(Inches(.9), Inches(1.35), Inches(10.8), Inches(1.25))
            title_style = title_box.text_frame.paragraphs[0]
            title_style.text = title; title_style.font.name = "Aptos Display"; title_style.font.size = Pt(42)
            title_style.font.bold = True; title_style.font.color.rgb = RGBColor(255, 255, 255)
            subtitle_box = slide.shapes.add_textbox(Inches(.95), Inches(2.8), Inches(10.5), Inches(1.0))
            p = subtitle_box.text_frame.paragraphs[0]; p.text = subtitle; p.font.size = Pt(19); p.font.color.rgb = RGBColor(194, 218, 231)
            y = 4.0
        else:
            line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.75), Inches(.55), Inches(.15), Inches(.55))
            line.fill.solid(); line.fill.fore_color.rgb = RGBColor(23, 123, 174); line.line.fill.background()
            title_box = slide.shapes.add_textbox(Inches(1.02), Inches(.48), Inches(11.4), Inches(.65))
            p = title_box.text_frame.paragraphs[0]; p.text = title; p.font.name = "Aptos Display"; p.font.size = Pt(28)
            p.font.bold = True; p.font.color.rgb = RGBColor(24, 52, 72)
            sub = slide.shapes.add_textbox(Inches(1.02), Inches(1.25), Inches(11.2), Inches(.62))
            p = sub.text_frame.paragraphs[0]; p.text = subtitle; p.font.size = Pt(14); p.font.color.rgb = RGBColor(92, 119, 139)
            rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.02), Inches(2.0), Inches(11.2), Inches(.015))
            rule.fill.solid(); rule.fill.fore_color.rgb = RGBColor(220, 230, 237); rule.line.fill.background()
            y = 2.35
        has_dashboard_screenshot = index == 11 and screenshot.exists()
        bullet_width = 4.95 if has_dashboard_screenshot else 11.2
        row_height = min(.79, 4.45 / max(1, len(bullets)))
        for bullet in bullets:
            box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.02), Inches(y), Inches(bullet_width), Inches(row_height - .09))
            box.fill.solid(); box.fill.fore_color.rgb = RGBColor(34, 72, 96) if index == 0 else RGBColor(255, 255, 255)
            box.line.color.rgb = RGBColor(62, 103, 126) if index == 0 else RGBColor(226, 234, 239)
            tf = box.text_frame; tf.clear(); tf.margin_left = Inches(.18); tf.margin_right = Inches(.16)
            tf.margin_top = Inches(.06); tf.margin_bottom = Inches(.03); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            paragraph = tf.paragraphs[0]; paragraph.text = bullet; paragraph.font.name = "Aptos"
            paragraph.font.size = Pt(15 if len(bullet) < 105 else 12); paragraph.font.color.rgb = RGBColor(246, 251, 253) if index == 0 else RGBColor(49, 76, 96)
            paragraph.level = 0
            y += row_height
        if has_dashboard_screenshot:
            slide.shapes.add_picture(str(screenshot), Inches(6.35), Inches(2.35), width=Inches(5.85))
            caption = slide.shapes.add_textbox(Inches(6.35), Inches(6.22), Inches(5.85), Inches(.3))
            p = caption.text_frame.paragraphs[0]
            p.text = "Actual prototype UI · synthetic demo records · local capture"
            p.font.size = Pt(9); p.font.color.rgb = RGBColor(92, 119, 139)
        foot = slide.shapes.add_textbox(Inches(.85), Inches(7.14), Inches(11.6), Inches(.2))
        p = foot.text_frame.paragraphs[0]; p.text = f"GLUCOTWIN AI  ·  RESEARCH PROTOTYPE  ·  SYNTHETIC DATA  ·  {index + 1:02d} / {len(slides):02d}"
        p.font.size = Pt(8); p.font.color.rgb = RGBColor(152, 173, 187) if index == 0 else RGBColor(133, 153, 166)
    presentation.save(DOCS / "presentation.pptx")


def add_presentation_pdf(slides: list[tuple[str, str, list[str]]]) -> None:
    page = (13.333 * inch, 7.5 * inch)
    document = SimpleDocTemplate(str(DOCS / "presentation.pdf"), pagesize=page,
                                 leftMargin=.85 * inch, rightMargin=.85 * inch,
                                 topMargin=.8 * inch, bottomMargin=.55 * inch)
    title_style = ParagraphStyle("SlideTitle", fontName="Helvetica-Bold", fontSize=25, leading=31, textColor=colors.HexColor("#183448"), spaceAfter=12)
    subtitle_style = ParagraphStyle("SlideSubtitle", fontName="Helvetica", fontSize=12, leading=17, textColor=colors.HexColor("#58758a"), spaceAfter=25)
    bullet_style = ParagraphStyle("SlideBullet", fontName="Helvetica", fontSize=14, leading=20, textColor=colors.HexColor("#334e62"), leftIndent=15, firstLineIndent=-11, spaceAfter=11)
    story = []
    screenshot = DOCS / "dashboard_synthetic.png"
    for slide_index, (title, subtitle, bullets) in enumerate(slides):
        story.append(Spacer(1, .55 * inch))
        story.append(Paragraph(title.replace("&", "&amp;"), title_style))
        story.append(Paragraph(subtitle.replace("&", "&amp;"), subtitle_style))
        if slide_index == 11 and screenshot.exists():
            left = [Paragraph("• " + bullet.replace("&", "&amp;"), bullet_style) for bullet in bullets]
            picture = Image(str(screenshot), width=5.65 * inch, height=3.53 * inch)
            table = Table([[left, picture]], colWidths=[4.25 * inch, 6.0 * inch], hAlign="LEFT")
            table.setStyle(TableStyle([
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ]))
            story.append(table)
        else:
            for bullet in bullets:
                story.append(Paragraph("• " + bullet.replace("&", "&amp;"), bullet_style))
        story.append(PageBreak())

    def decorate(canvas, doc):
        canvas.saveState()
        width, height = page
        canvas.setFillColor(colors.HexColor("#f4f8fb")); canvas.rect(0, 0, width, height, fill=1, stroke=0)
        if doc.page == 1:
            canvas.setFillColor(colors.HexColor("#1e909d")); canvas.circle(width * .88, height * .87, .78 * inch, fill=1, stroke=0)
            canvas.setFillColor(colors.HexColor("#143148")); canvas.rect(0, height - .17 * inch, width, .17 * inch, fill=1, stroke=0)
        else:
            canvas.setStrokeColor(colors.HexColor("#dce6ed")); canvas.line(.85 * inch, .42 * inch, width - .85 * inch, .42 * inch)
            canvas.setFillColor(colors.HexColor("#8498a8")); canvas.setFont("Helvetica", 7)
            canvas.drawString(.85 * inch, .24 * inch, "GLUCOTWIN AI · SYNTHETIC DATA · NOT MEDICAL ADVICE")
            canvas.drawRightString(width - .85 * inch, .24 * inch, f"{doc.page:02d} / {len(slides):02d}")
        canvas.restoreState()

    document.build(story, onFirstPage=decorate, onLaterPages=decorate)


def draw_architecture_pdf() -> None:
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import landscape, A3

    output = DOCS / "architecture_diagram.pdf"
    width, height = landscape(A3)
    pdf = canvas.Canvas(str(output), pagesize=(width, height))
    navy = colors.HexColor("#18354d"); blue = colors.HexColor("#eaf3fb"); teal = colors.HexColor("#e8f6f3")
    edge = colors.HexColor("#dce6ed"); ink = colors.HexColor("#21445e"); muted = colors.HexColor("#668096")
    pdf.setFillColor(colors.HexColor("#f4f8fb")); pdf.rect(0, 0, width, height, fill=1, stroke=0)
    pdf.setFillColor(navy); pdf.setFont("Helvetica-Bold", 24); pdf.drawString(35, height - 46, "GlucoTwin AI 2.0 · System Architecture")
    pdf.setFillColor(muted); pdf.setFont("Helvetica", 11); pdf.drawString(36, height - 66, "Synthetic-data patient twin · experimental 120-minute estimate · not for clinical use")

    def box(x, y, w, h, title, lines, fill=colors.white, title_color=ink):
        pdf.setFillColor(fill); pdf.setStrokeColor(edge); pdf.setLineWidth(1.2); pdf.roundRect(x, y, w, h, 10, fill=1, stroke=1)
        pdf.setFillColor(title_color); pdf.setFont("Helvetica-Bold", 9); pdf.drawCentredString(x + w / 2, y + h - 22, title)
        pdf.setFillColor(muted); pdf.setFont("Helvetica", 7.2)
        for idx, line in enumerate(lines):
            pdf.drawCentredString(x + w / 2, y + h - 43 - 14 * idx, line)

    def arrow(x1, y1, x2, y2):
        pdf.setStrokeColor(colors.HexColor("#628099")); pdf.setFillColor(colors.HexColor("#628099")); pdf.setLineWidth(1.7)
        pdf.line(x1, y1, x2, y2)
        import math
        angle = math.atan2(y2 - y1, x2 - x1); size = 7
        for turn in (-.55, .55):
            pdf.line(x2, y2, x2 - size * math.cos(angle + turn), y2 - size * math.sin(angle + turn))

    margin, gap = 22, 5
    box_w = (width - 2 * margin - 8 * gap) / 9
    y, h = height - 250, 130
    main = [
        ("Data streams", ["Historical EHR", "Wearable / IoT", "Synthetic sources"], colors.white),
        ("Validation", ["Pydantic checks", "Units + ranges", "Timezone / duplicate"], blue),
        ("Patient twin", ["Profile + history", "Latest observations", "Updateable state"], teal),
        ("Temporal features", ["Past-only lags", "Rolling summaries", "Training medians"], blue),
        ("ML prediction", ["Saved model pipeline", "120-minute event", "Uncalibrated score"], colors.HexColor("#eaf2f8")),
        ("Explainability", ["Observed signals", "Global feature ranks", "Quality warnings"], teal),
        ("FastAPI", ["Validated JSON routes", "Health + /docs", "Exact-origin CORS"], colors.white),
        ("SQLite", ["Patient + readings", "Prediction history", "Local persistence"], blue),
        ("Dashboard", ["Responsive UI", "Measured glucose trend", "Results + safety notes"], colors.white),
    ]
    pdf.setFillColor(muted); pdf.setFont("Helvetica-Bold", 9); pdf.drawString(margin, y + h + 22, "ONLINE DATA, PREDICTION, AND PRESENTATION FLOW")
    xs = []
    for index, (title, lines, fill) in enumerate(main):
        x = margin + index * (box_w + gap); xs.append(x)
        box(x, y, box_w, h, title, lines, fill)
        if index:
            arrow(xs[index - 1] + box_w + 1, y + h / 2, x - 2, y + h / 2)

    train_y = 196
    pdf.setFillColor(muted); pdf.setFont("Helvetica-Bold", 9); pdf.drawString(margin, train_y + 112, "OFFLINE TRAINING / EVALUATION · EXPLICIT COMMAND")
    train_gap = 10; train_w = (width - 2 * margin - 4 * train_gap) / 5
    training = [
        ("Synthetic generator", ["36 profiles · 14 days", "15-minute rows", "Fixed seed + marker"]),
        ("Label + features", ["Past-only inputs", "Next 8-row event label", "Incomplete windows removed"]),
        ("Patient split", ["Train / validation / test", "No person crosses groups", "Fixed split seed"]),
        ("Model comparison", ["Prior / logistic baseline", "Random Forest candidate", "Select on validation PR-AUC"]),
        ("Saved artifacts", ["joblib inference bundle", "Test metrics + confusion", "Synthetic limitations report"]),
    ]
    train_xs = []
    for index, (title, lines) in enumerate(training):
        x = margin + index * (train_w + train_gap); train_xs.append(x)
        box(x, train_y, train_w, 88, title, lines, blue if index < 4 else teal)
        if index:
            arrow(train_xs[index - 1] + train_w + 2, train_y + 44, x - 5, train_y + 44)
    pdf.setFillColor(muted); pdf.setFont("Helvetica", 7.5)
    pdf.drawString(margin, train_y - 14, "The generated synthetic-trained pipeline is loaded by the prediction route; training never implies clinical validity.")

    security_x, security_y, security_w, security_h = margin, 50, width - 2 * margin, 56
    pdf.setStrokeColor(colors.HexColor("#80a3bd")); pdf.setDash(5, 4); pdf.roundRect(security_x, security_y, security_w, security_h, 10, fill=0, stroke=1); pdf.setDash()
    pdf.setFillColor(ink); pdf.setFont("Helvetica-Bold", 8); pdf.drawString(security_x + 12, security_y + 36, "PROTOTYPE SAFEGUARDS")
    pdf.setFillColor(muted); pdf.setFont("Helvetica", 8)
    pdf.drawString(security_x + 160, security_y + 36, "Synthetic-labelled records · parameterized SQL · exact-origin CORS · environment configuration")
    pdf.drawString(security_x + 160, security_y + 19, "No authentication, production controls, clinical validation, regulatory clearance, or care use")
    pdf.setFillColor(muted); pdf.setFont("Helvetica", 7)
    pdf.drawString(margin, 12, "GlucoTwin AI 2.0 · Editable SVG source: docs/architecture_diagram.svg · Synthetic data only")
    pdf.showPage(); pdf.save()


def main() -> None:
    DOCS.mkdir(parents=True, exist_ok=True)
    slides = slides_data()
    add_powerpoint(slides)
    add_presentation_pdf(slides)
    draw_architecture_pdf()
    print(f"Generated {len(slides)}-slide presentation (PPTX/PDF) and architecture PDF in {DOCS}.")


if __name__ == "__main__":
    main()
